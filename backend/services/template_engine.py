"""
ROI Dashboard — Premium PowerPoint Template Engine
One dedicated layout per slide type. No generic fallback abuse.
"""
from __future__ import annotations

import io
import os
import re
import uuid

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

from models.schemas import ReportOutline, SlideOutline
from config import settings

# ── Palette ───────────────────────────────────────────────────────────────────

# Databricks web-inspired core palette
BRAND_RED    = RGBColor(0xFF, 0x36, 0x21)
BRAND_TEAL   = RGBColor(0x1B, 0x51, 0x62)
BRAND_GREEN  = RGBColor(0x00, 0xA9, 0x72)
BRAND_ORANGE = RGBColor(0xFF, 0xAB, 0x00)
BRAND_PURPLE = RGBColor(0xAB, 0x40, 0x57)

HDR_BG    = RGBColor(0x1B, 0x31, 0x39)
COVER_BG  = RGBColor(0x14, 0x3D, 0x4A)
SLIDE_BG  = RGBColor(0xF9, 0xF7, 0xF4)
CARD_BG   = RGBColor(0xFF, 0xFF, 0xFF)
CARD_BOR  = RGBColor(0xDC, 0xE0, 0xE2)
SURFACE   = RGBColor(0xEE, 0xED, 0xE9)
TXT_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TXT_DARK  = RGBColor(0x1B, 0x31, 0x39)
TXT_MID   = RGBColor(0x5A, 0x6F, 0x77)
TXT_LIGHT = RGBColor(0x90, 0xA5, 0xB1)
DIVIDER   = RGBColor(0xDC, 0xE0, 0xE2)

ACCENTS = [BRAND_RED, BRAND_TEAL, BRAND_GREEN, BRAND_ORANGE, BRAND_PURPLE]
CHART_HEX = ["#FF3621", "#1B5162", "#00A972", "#FFAB00", "#AB4057", "#618794", "#90A5B1"]

# ── Geometry ──────────────────────────────────────────────────────────────────

W = Inches(13.33)
H = Inches(7.50)

STRIPE_H = Inches(0.07)
HDR_H    = Inches(0.78)
HDR_BOT  = STRIPE_H + HDR_H           # 0.85"

FOOT_H   = Inches(0.38)
FOOT_Y   = H - FOOT_H                 # 7.12"

MG    = Inches(0.38)
CTOP  = HDR_BOT + Inches(0.18)        # 1.03"
CBOT  = FOOT_Y - Inches(0.10)         # 7.02"
CAVAIL = CBOT - CTOP                   # 5.99"
CW    = W - 2 * MG                    # 12.57"

# ── Primitives ────────────────────────────────────────────────────────────────

def _bg(slide, color: RGBColor = SLIDE_BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _rect(slide, l, t, w, h, fill: RGBColor,
          border: RGBColor = None, bpt: float = 0.75):
    shp = slide.shapes.add_shape(1, l, t, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if border:
        shp.line.color.rgb = border
        shp.line.width = Pt(bpt)
    else:
        shp.line.fill.background()
    return shp


def _rrect(slide, l, t, w, h, fill: RGBColor,
           border: RGBColor = None, bpt: float = 0.75, adj: int = 5000):
    shp = slide.shapes.add_shape(1, l, t, w, h)
    try:
        sp_pr = shp._element.spPr
        pst = sp_pr.find(qn("a:prstGeom"))
        if pst is not None:
            pst.set("prst", "roundRect")
            av = pst.find(qn("a:avLst"))
            if av is None:
                av = etree.SubElement(pst, qn("a:avLst"))
            for gd in list(av.findall(qn("a:gd"))):
                av.remove(gd)
            gd = etree.SubElement(av, qn("a:gd"))
            gd.set("name", "adj")
            gd.set("fmla", f"val {adj}")
    except Exception:
        pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if border:
        shp.line.color.rgb = border
        shp.line.width = Pt(bpt)
    else:
        shp.line.fill.background()
    return shp


def _txt(slide, text: str, l, t, w, h,
         size: float = 11, bold: bool = False, italic: bool = False,
         color: RGBColor = TXT_DARK, align: PP_ALIGN = PP_ALIGN.LEFT,
         wrap: bool = True, font: str = "DM Sans",
         ls: float = 1.15, anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
         space_after: float = 1):
    txb = slide.shapes.add_textbox(l, t, w, h)
    tf = txb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(0)
    tf.margin_right = Pt(0)
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = ls
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = str(text) if text is not None else ""
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    return txb


def _pic(slide, img: bytes, l, t, w, h):
    if img:
        slide.shapes.add_picture(io.BytesIO(img), l, t, w, h)


def _clean(s: str) -> str:
    if not s:
        return ""
    s = re.sub(r"\*\*(.*?)\*\*", r"\1", s)
    s = re.sub(r"\*(.*?)\*", r"\1", s)
    return s.strip()


def _fit(text: str, n: int) -> str:
    t = _clean(text)
    return t if len(t) <= n else t[:max(0, n - 1)].rstrip() + "…"

# ── Header / Footer ───────────────────────────────────────────────────────────

def _header(slide, title: str, num: int, total: int, accent: RGBColor = BRAND_RED):
    _rect(slide, 0, 0, W, STRIPE_H, accent)
    _rect(slide, 0, STRIPE_H, W, HDR_H, HDR_BG)
    # Subtle right-side color block in header
    _rect(slide, W - Inches(3.2), STRIPE_H, Inches(3.2), HDR_H,
          RGBColor(0x1C, 0x2C, 0x50))
    _txt(slide, title,
         MG, STRIPE_H + Inches(0.12),
         W - MG - Inches(3.4), HDR_H - Inches(0.16),
         size=20, bold=True, color=TXT_WHITE, wrap=False)
    # Slide counter pill
    pill_w, pill_h = Inches(1.30), Inches(0.34)
    pill_l = W - MG - pill_w
    pill_t = STRIPE_H + (HDR_H - pill_h) / 2
    _rrect(slide, pill_l, pill_t, pill_w, pill_h,
           RGBColor(0x2A, 0x3A, 0x60), adj=16000)
    _txt(slide, f"{num} / {total}",
         pill_l, pill_t, pill_w, pill_h,
         size=9, bold=True, color=RGBColor(0x8A, 0xA8, 0xD0),
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def _footer(slide, event_name: str):
    _rect(slide, 0, FOOT_Y, W, Inches(0.012), DIVIDER)
    _rect(slide, 0, FOOT_Y + Inches(0.012), W, FOOT_H - Inches(0.012), SLIDE_BG)
    _txt(slide, _fit(event_name, 65),
         MG, FOOT_Y + Inches(0.08), Inches(7), Inches(0.24),
         size=8.5, italic=True, color=TXT_LIGHT)
    _txt(slide, "Powered by Gemini AI  ·  Databricks  ·  LatentView",
         W - Inches(4.8), FOOT_Y + Inches(0.08), Inches(4.45), Inches(0.24),
         size=8.5, color=TXT_LIGHT, align=PP_ALIGN.RIGHT)

# ── Shared Components ─────────────────────────────────────────────────────────

def _kpi_card(slide, l, t, w, h, value: str, label: str,
              accent: RGBColor = BRAND_RED, sub: str = ""):
    _rrect(slide, l, t, w, h, CARD_BG, border=CARD_BOR, bpt=0.5, adj=3500)
    _rect(slide, l, t, w, Inches(0.055), accent)
    _txt(slide, _fit(value, 16),
         l + Inches(0.10), t + Inches(0.10),
         w - Inches(0.20), Inches(0.76),
         size=26, bold=True, color=TXT_DARK,
         align=PP_ALIGN.CENTER, wrap=False)
    _txt(slide, label.upper(),
         l + Inches(0.08), t + Inches(0.90),
         w - Inches(0.16), Inches(0.26),
         size=7.5, bold=True, color=accent, align=PP_ALIGN.CENTER)
    if sub:
        _txt(slide, _fit(sub, 38),
             l + Inches(0.08), t + Inches(1.18),
             w - Inches(0.16), Inches(0.20),
             size=7, color=TXT_LIGHT, align=PP_ALIGN.CENTER)


def _bullets(slide, points: list[str], l, t, w, max_h,
             fs: float = 10.5, accent: RGBColor = BRAND_RED):
    y = t
    row_h = Inches(0.38)
    for pt in points:
        if y + row_h > t + max_h:
            break
        text = _fit(pt, 130)
        if not text:
            continue
        _rect(slide, l, y + Inches(0.14), Inches(0.065), Inches(0.065), accent)
        _txt(slide, text,
             l + Inches(0.17), y, w - Inches(0.20), row_h,
             size=fs, color=TXT_MID, wrap=True, ls=1.10, space_after=0)
        y += row_h
    return y


def _insight_box(slide, text: str, l, t, w, h=None):
    if not text:
        return
    bh = h or Inches(1.08)
    _rrect(slide, l, t, w, bh,
           RGBColor(0xFF, 0xF7, 0xCC),
           border=RGBColor(0xF5, 0xC5, 0x42), bpt=1.0, adj=2500)
    _rect(slide, l, t, Inches(0.07), bh, RGBColor(0xF5, 0xC5, 0x42))
    _txt(slide, "KEY INSIGHT",
         l + Inches(0.18), t + Inches(0.08),
         w - Inches(0.28), Inches(0.22),
         size=7.5, bold=True, color=RGBColor(0x78, 0x49, 0x00))
    _txt(slide, _fit(text, 220),
         l + Inches(0.18), t + Inches(0.30),
         w - Inches(0.28), bh - Inches(0.36),
         size=10.5, color=TXT_DARK, wrap=True, ls=1.12)


def _metric_panel(slide, l, t, w, h,
                  headline: str, label: str, sub: str = "",
                  accent: RGBColor = BRAND_RED):
    _rrect(slide, l, t, w, h, HDR_BG, adj=2500)
    _rect(slide, l, t, Inches(0.10), h, accent)
    _txt(slide, _fit(headline, 14),
         l + Inches(0.22), t + Inches(0.40),
         w - Inches(0.36), Inches(1.35),
         size=36, bold=True, color=TXT_WHITE,
         align=PP_ALIGN.CENTER, wrap=False)
    _txt(slide, label.upper(),
         l + Inches(0.18), t + Inches(1.84),
         w - Inches(0.32), Inches(0.30),
         size=8.5, bold=True, color=accent, align=PP_ALIGN.CENTER)
    if sub:
        _txt(slide, _fit(sub, 110),
             l + Inches(0.18), t + Inches(2.22),
             w - Inches(0.32), h - Inches(2.34),
             size=8.5, color=RGBColor(0x8A, 0xA8, 0xC8),
             align=PP_ALIGN.CENTER, wrap=True, ls=1.12)


def _section_card(slide, l, t, w, h, title: str, items: list[str],
                  accent: RGBColor = BRAND_TEAL, bg: RGBColor = CARD_BG):
    _rrect(slide, l, t, w, h, bg, border=CARD_BOR, bpt=0.5, adj=2500)
    _rect(slide, l, t, w, Inches(0.042), accent)
    _txt(slide, title.upper(),
         l + Inches(0.14), t + Inches(0.09),
         w - Inches(0.28), Inches(0.24),
         size=7.5, bold=True, color=accent)
    y = t + Inches(0.38)
    for item in items:
        if y + Inches(0.32) > t + h - Inches(0.06):
            break
        _txt(slide, _fit(f"› {item}", 88),
             l + Inches(0.12), y, w - Inches(0.24), Inches(0.32),
             size=9.5, color=TXT_MID, wrap=True, ls=1.10, space_after=0)
        y += Inches(0.32)


def _action_card(slide, l, t, w, h, number: int, text: str, accent: RGBColor):
    _rrect(slide, l, t, w, h, CARD_BG, border=accent, bpt=1.0, adj=3000)
    bs = Inches(0.36)
    bx = l + Inches(0.12)
    by = t + (h - bs) / 2
    _rrect(slide, bx, by, bs, bs, accent, adj=16000)
    _txt(slide, str(number), bx, by, bs, bs,
         size=13, bold=True, color=TXT_WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    _txt(slide, _fit(text, 110),
         l + Inches(0.60), t + Inches(0.10),
         w - Inches(0.74), h - Inches(0.20),
         size=9.5, color=TXT_DARK, wrap=True, ls=1.12)

# ── Chart Generators ──────────────────────────────────────────────────────────

def _chart_bar(labels: list[str], values: list[float],
               w_in: float = 5.5, h_in: float = 4.0,
               colors: list[str] = None) -> bytes:
    n = len(labels)
    if n == 0:
        return b""
    fig_h = max(2.0, min(h_in, n * 0.58 + 0.5))
    fig, ax = plt.subplots(figsize=(w_in, fig_h))
    fig.patch.set_facecolor("#FFFFFF")
    ax.set_facecolor("#FFFFFF")

    bar_cols = colors or [CHART_HEX[i % len(CHART_HEX)] for i in range(n)]
    ax.barh(range(n), values, color=bar_cols, height=0.56, edgecolor="none", zorder=3)

    max_v = max(values) if values else 1
    for i, (v, lbl) in enumerate(zip(values, labels)):
        ax.text(-max_v * 0.025, i, _fit(lbl, 28),
                va="center", ha="right", color="#4A5568", fontsize=9, fontname="DM Sans")
        fmt = f"${int(v):,}" if max_v > 1000 else (f"{v:.1f}%" if max_v < 101 else f"{int(v):,}")
        ax.text(v + max_v * 0.025, i, fmt,
                va="center", ha="left", color="#1A202C",
                fontsize=9, fontname="DM Sans", fontweight="bold")

    ax.set_xlim(-max_v * 0.40, max_v * 1.30)
    ax.set_ylim(-0.5, n - 0.5)
    ax.invert_yaxis()
    for sp in ax.spines.values():
        sp.set_visible(False)
    ax.tick_params(length=0)
    ax.set_xticks([])
    ax.set_yticks([])
    for xi in np.linspace(0, max_v, 4)[1:]:
        ax.axvline(xi, color="#E1E7F0", linewidth=0.8, zorder=1)

    plt.tight_layout(pad=0.25)
    buf = io.BytesIO()
    plt.savefig(buf, format="png", dpi=150, bbox_inches="tight",
                facecolor="#FFFFFF", edgecolor="none")
    plt.close(fig)
    buf.seek(0)
    return buf.read()


def _chart_funnel(stages: list[str], values: list[float],
                  w_in: float = 5.5, h_in: float = 4.5) -> bytes:
    n = len(stages)
    if n < 2:
        return b""
    fig, ax = plt.subplots(figsize=(w_in, h_in))
    fig.patch.set_facecolor("#FFFFFF")
    ax.set_facecolor("#FFFFFF")

    max_v = values[0] if values else 1
    seg_h = 1.0 / n
    stage_cols = ["#FF3621", "#FF6B35", "#FF9F45", "#00B86A", "#0094BD", "#7C3AED"]

    for i in range(n):
        pct = values[i] / max_v if max_v > 0 else 0
        prev_pct = values[i - 1] / max_v if (i > 0 and max_v > 0) else 1.0
        y = 1.0 - (i + 1) * seg_h
        col = stage_cols[i % len(stage_cols)]
        xs = [(1 - prev_pct) / 2, (1 + prev_pct) / 2,
              (1 + pct) / 2,      (1 - pct) / 2]
        ys = [y + seg_h, y + seg_h, y, y]
        ax.fill(xs, ys, color=col, alpha=0.90, zorder=2)
        ax.plot(xs + [xs[0]], ys + [ys[0]], color="#FFFFFF", lw=1.8, zorder=3)

        lbl = _fit(stages[i], 20)
        cnt = f"{int(values[i]):,}" if i < len(values) else ""
        conv = f"  ({pct * 100:.0f}%)" if i > 0 else "  (100%)"
        mid_y = y + seg_h / 2
        ax.text(0.5, mid_y + 0.018, lbl,
                ha="center", va="center", color="#FFFFFF",
                fontsize=9, fontname="DM Sans", fontweight="bold", zorder=4)
        ax.text(0.5, mid_y - 0.022, f"{cnt}{conv}",
                ha="center", va="center", color="#FFFFFF",
                fontsize=7.5, fontname="DM Sans", zorder=4)

    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    plt.tight_layout(pad=0.15)
    buf = io.BytesIO()
    plt.savefig(buf, format="png", dpi=150, bbox_inches="tight",
                facecolor="#FFFFFF", edgecolor="none")
    plt.close(fig)
    buf.seek(0)
    return buf.read()

# ── Data Extraction ───────────────────────────────────────────────────────────

_NUM_RE = re.compile(r"(\$?[\d,]+(?:\.\d+)?[kKmMbBxX%]?)")


def _extract_kpi_pairs(points: list[str]) -> list[tuple[str, str]]:
    """Return [(label, value), ...] for up to 4 KPI cards."""
    pairs = []
    for pt in points:
        t = _clean(pt)
        m = re.match(r"^(.{4,40}?)[:\–\—]\s*(.{1,22})$", t)
        if m:
            lbl, val = m.group(1).strip()[:38], m.group(2).strip()[:18]
            if _NUM_RE.search(val):
                pairs.append((lbl, val))
                continue
        mv = _NUM_RE.search(t)
        if mv:
            val = mv.group(1)
            lbl = t.replace(mv.group(0), "").strip(":-– ")[:38] or "Metric"
            pairs.append((lbl, val))
    return pairs[:4]


def _extract_bar_data(points: list[str]) -> tuple[list[str], list[float]]:
    labels, values = [], []
    for pt in points:
        t = _clean(pt)
        m = re.search(r"(.{4,38}?)[:\–\—]\s*(?:\$?)([\d,\.]+)", t)
        if m:
            try:
                v = float(m.group(2).replace(",", ""))
                if v > 0:
                    labels.append(m.group(1).strip()[:28])
                    values.append(v)
            except ValueError:
                pass
    return labels[:8], values[:8]


def _extract_funnel_data(points: list[str]) -> tuple[list[str], list[float]]:
    stages, values = [], []
    for pt in points:
        t = _clean(pt)
        m = re.search(r"(.{3,28}?)[:\–\—]\s*\*{0,2}([\d,]+)", t)
        if m:
            try:
                stages.append(m.group(1).strip()[:24])
                values.append(float(m.group(2).replace(",", "")))
            except ValueError:
                pass
    return stages, values


def _headline(points: list[str]) -> tuple[str, str, str]:
    if not points:
        return "—", "Key Metric", ""
    t = _clean(points[0])
    mv = _NUM_RE.search(t)
    if mv:
        val = mv.group(1)
        lbl = t.replace(mv.group(0), "").strip(":-– ")[:32] or "Headline"
        sub = _fit(" ".join(_clean(p) for p in points[1:3]), 110)
        return val, lbl, sub
    parts = re.split(r"[:\–\—]", t, maxsplit=1)
    if len(parts) == 2:
        return _fit(parts[1], 16), _fit(parts[0], 32), _fit(" ".join(_clean(p) for p in points[1:3]), 110)
    return t[:16], "Metric", ""

# ── Slide Type Detection ──────────────────────────────────────────────────────

_SLIDE_MAP = {
    1: "exec",
    2: "pipeline",
    3: "funnel",
    4: "segment",
    5: "platform",
    6: "financial",
    7: "actions",
}


def _stype(title: str, has_table: bool, num: int) -> str:
    if num in _SLIDE_MAP:
        return _SLIDE_MAP[num]
    t = title.lower()
    if has_table:
        return "table"
    if any(w in t for w in ("exec", "summary", "overview", "snapshot")):
        return "exec"
    if any(w in t for w in ("pipeline", "revenue")):
        return "pipeline"
    if any(w in t for w in ("funnel", "stage", "conversion", "health")):
        return "funnel"
    if any(w in t for w in ("segment", "audience", "region", "breakdown")):
        return "segment"
    if any(w in t for w in ("platform", "data quality", "gap")):
        return "platform"
    if any(w in t for w in ("financial", "roi", "cost", "spend", "efficien")):
        return "financial"
    if any(w in t for w in ("risk", "action", "next step", "recommend")):
        return "actions"
    return "standard"

# ── Slide 1: Executive Summary ────────────────────────────────────────────────

def _slide_exec(slide, data: SlideOutline, total: int, event_name: str):
    _header(slide, data.title, data.number, total, BRAND_RED)
    _footer(slide, event_name)

    kpis = _extract_kpi_pairs(data.points)
    n = min(len(kpis), 4) or 4
    gap = Inches(0.16)
    cw = (CW - (n - 1) * gap) / n
    ch = Inches(1.48)

    if kpis:
        for i, (lbl, val) in enumerate(kpis[:4]):
            _kpi_card(slide, MG + i * (cw + gap), CTOP, cw, ch,
                      val, lbl, ACCENTS[i % 4])
    else:
        for i, pt in enumerate(data.points[:4]):
            t = _clean(pt)
            parts = re.split(r"[:\–\—]", t, 1)
            val = parts[1].strip()[:16] if len(parts) > 1 else t[:14]
            lbl = parts[0].strip()[:32] if len(parts) > 1 else f"Metric {i+1}"
            _kpi_card(slide, MG + i * (cw + gap), CTOP, cw, ch,
                      val, lbl, ACCENTS[i % 4])

    rest = data.points[min(4, len(data.points)):]
    bt = CTOP + ch + Inches(0.16)
    ins_h = Inches(1.08) if data.insight else Inches(0)
    bh = CBOT - bt - ins_h - Inches(0.08)

    if rest:
        half = (len(rest) + 1) // 2
        col_w = (CW - Inches(0.28)) / 2
        _bullets(slide, rest[:half], MG, bt, col_w, bh, fs=10, accent=BRAND_RED)
        _bullets(slide, rest[half:], MG + col_w + Inches(0.28), bt, col_w, bh, fs=10, accent=BRAND_RED)

    if data.insight:
        _insight_box(slide, data.insight, MG, CBOT - ins_h, CW)

# ── Slide 2: Pipeline & Revenue ───────────────────────────────────────────────

def _slide_pipeline(slide, data: SlideOutline, total: int, event_name: str):
    _header(slide, data.title, data.number, total, BRAND_TEAL)
    _footer(slide, event_name)

    lw = Inches(3.80)
    rx = MG + lw + Inches(0.22)
    rw = W - rx - MG

    headline, lbl, sub = _headline(data.points)
    _metric_panel(slide, MG, CTOP, lw, CAVAIL, headline, lbl, sub, BRAND_TEAL)

    ins_h = Inches(1.08) if data.insight else Inches(0)
    labels, values = _extract_bar_data(data.points)

    if len(values) >= 2:
        ch = CAVAIL * 0.54
        try:
            img = _chart_bar(labels, values, w_in=rw.inches - 0.12, h_in=ch.inches - 0.10)
            _rrect(slide, rx, CTOP, rw, ch, CARD_BG, border=CARD_BOR, bpt=0.4, adj=2500)
            _txt(slide, "BREAKDOWN BY SEGMENT",
                 rx + Inches(0.15), CTOP + Inches(0.07),
                 rw - Inches(0.3), Inches(0.22),
                 size=7.5, bold=True, color=BRAND_TEAL)
            _pic(slide, img, rx + Inches(0.08), CTOP + Inches(0.30),
                 rw - Inches(0.16), ch - Inches(0.38))
        except Exception:
            ch = Inches(0)
        bt = CTOP + ch + Inches(0.14)
        bh = CBOT - bt - ins_h - Inches(0.08)
        _bullets(slide, data.points, rx, bt, rw, bh, fs=10, accent=BRAND_TEAL)
    else:
        bh = CAVAIL - ins_h - Inches(0.05)
        _bullets(slide, data.points, rx, CTOP, rw, bh, fs=10.5, accent=BRAND_TEAL)

    if data.insight:
        _insight_box(slide, data.insight, rx, CBOT - ins_h, rw)

# ── Slide 3: Funnel & Conversion ──────────────────────────────────────────────

def _slide_funnel(slide, data: SlideOutline, total: int, event_name: str):
    _header(slide, data.title, data.number, total, BRAND_ORANGE)
    _footer(slide, event_name)

    fw = Inches(5.80)
    rx = MG + fw + Inches(0.22)
    rw = W - rx - MG

    stages, values = _extract_funnel_data(data.points)

    if len(stages) >= 2:
        try:
            img = _chart_funnel(stages, values,
                                w_in=fw.inches - 0.12, h_in=CAVAIL.inches - 0.18)
            _rrect(slide, MG, CTOP, fw, CAVAIL, CARD_BG, border=CARD_BOR, bpt=0.4, adj=2500)
            _txt(slide, "CONVERSION FUNNEL",
                 MG + Inches(0.15), CTOP + Inches(0.07),
                 fw - Inches(0.3), Inches(0.22),
                 size=7.5, bold=True, color=BRAND_ORANGE)
            _pic(slide, img, MG + Inches(0.08), CTOP + Inches(0.30),
                 fw - Inches(0.16), CAVAIL - Inches(0.38))
        except Exception:
            pass

        # Stage stat cards on right
        n_s = min(len(stages), 5)
        card_gap = Inches(0.10)
        card_h = (CAVAIL - (n_s - 1) * card_gap) / max(n_s, 1)
        scols = [BRAND_RED, BRAND_ORANGE, BRAND_GREEN, BRAND_TEAL, BRAND_PURPLE]

        for i in range(n_s):
            cy = CTOP + i * (card_h + card_gap)
            col = scols[i % len(scols)]
            _rrect(slide, rx, cy, rw, card_h, CARD_BG, border=col, bpt=0.9, adj=2500)
            _rect(slide, rx, cy, Inches(0.06), card_h, col)
            _txt(slide, _fit(stages[i], 26),
                 rx + Inches(0.14), cy + Inches(0.06),
                 rw - Inches(0.22), Inches(0.26),
                 size=8.5, bold=True, color=col)
            if i < len(values):
                cnt = f"{int(values[i]):,}"
                conv = f"  ({values[i] / values[0] * 100:.0f}%)" if i > 0 and values[0] > 0 else "  (100%)"
                _txt(slide, cnt + conv,
                     rx + Inches(0.14), cy + Inches(0.33),
                     rw - Inches(0.22), card_h - Inches(0.40),
                     size=13, bold=True, color=TXT_DARK)
    else:
        # Not enough parseable data: fall back to standard layout
        _slide_standard(slide, data, total, event_name)
        return

    if data.insight:
        # Narrow insight bar along bottom full-width
        _insight_box(slide, data.insight, MG, CBOT - Inches(0.92), CW, Inches(0.92))

# ── Slide 4: Audience & Segment ───────────────────────────────────────────────

def _slide_segment(slide, data: SlideOutline, total: int, event_name: str):
    _header(slide, data.title, data.number, total, BRAND_PURPLE)
    _footer(slide, event_name)

    cw_chart = Inches(7.00)
    rx = MG + cw_chart + Inches(0.22)
    rw = W - rx - MG

    labels, values = _extract_bar_data(data.points)

    if len(values) >= 2:
        pairs = sorted(zip(labels, values), key=lambda x: x[1], reverse=True)
        sl, sv = zip(*pairs[:8])
        cols = [CHART_HEX[i % len(CHART_HEX)] for i in range(len(sl))]
        try:
            img = _chart_bar(list(sl), list(sv), colors=cols,
                             w_in=cw_chart.inches - 0.14, h_in=CAVAIL.inches - 0.22)
            _rrect(slide, MG, CTOP, cw_chart, CAVAIL, CARD_BG, border=CARD_BOR, bpt=0.4, adj=2500)
            _txt(slide, "SEGMENT PERFORMANCE RANKING",
                 MG + Inches(0.15), CTOP + Inches(0.07),
                 cw_chart - Inches(0.3), Inches(0.22),
                 size=7.5, bold=True, color=BRAND_PURPLE)
            _pic(slide, img, MG + Inches(0.08), CTOP + Inches(0.30),
                 cw_chart - Inches(0.16), CAVAIL - Inches(0.38))
        except Exception:
            pass
    else:
        _bullets(slide, data.points, MG, CTOP,
                 cw_chart, CAVAIL, fs=10.5, accent=BRAND_PURPLE)

    # Right column: top performers card + watch areas card + insight
    top_h = CAVAIL * 0.44
    bot_h = CAVAIL * 0.38
    ins_h = CAVAIL - top_h - bot_h - Inches(0.20)

    _section_card(slide, rx, CTOP, rw, top_h,
                  "Top Performers", data.points[:3],
                  BRAND_GREEN, RGBColor(0xF0, 0xFB, 0xF4))
    _section_card(slide, rx, CTOP + top_h + Inches(0.10), rw, bot_h,
                  "Watch Areas",
                  data.points[3:6] or ["No underperforming segments identified"],
                  BRAND_ORANGE, RGBColor(0xFF, 0xF8, 0xF0))
    if data.insight and ins_h > Inches(0.60):
        _insight_box(slide, data.insight,
                     rx, CTOP + top_h + bot_h + Inches(0.20), rw, ins_h)

# ── Slide 5: Platform & Data Quality ─────────────────────────────────────────

def _slide_platform(slide, data: SlideOutline, total: int, event_name: str):
    _header(slide, data.title, data.number, total, BRAND_TEAL)
    _footer(slide, event_name)

    col_gap = Inches(0.20)
    col_w = (CW - 2 * col_gap) / 3

    card_h = CAVAIL - (Inches(1.08) if data.insight else Inches(0)) - Inches(0.12)

    cats = [
        ("Data Coverage",     BRAND_TEAL,   data.points[:3]),
        ("Query Performance", BRAND_PURPLE,  data.points[3:6]),
        ("Data Completeness", BRAND_ORANGE,  data.points[6:9]),
    ]
    for i, (title, col, pts) in enumerate(cats):
        cx = MG + i * (col_w + col_gap)
        items = pts if pts else ["No data available for this category"]
        _section_card(slide, cx, CTOP, col_w, card_h, title, items, col)

    if data.insight:
        _insight_box(slide, data.insight, MG, CBOT - Inches(1.08), CW)

# ── Slide 6: Financial Efficiency & ROI ──────────────────────────────────────

def _slide_financial(slide, data: SlideOutline, total: int, event_name: str):
    _header(slide, data.title, data.number, total, BRAND_GREEN)
    _footer(slide, event_name)

    # Top KPI cards row
    kpis = _extract_kpi_pairs(data.points)
    n = min(len(kpis), 3) or 3
    cgap = Inches(0.18)
    cw_card = (CW - (n - 1) * cgap) / max(n, 1)
    ch_card = Inches(1.44)

    if kpis:
        for i, (lbl, val) in enumerate(kpis[:3]):
            _kpi_card(slide, MG + i * (cw_card + cgap), CTOP,
                      cw_card, ch_card, val, lbl,
                      [BRAND_GREEN, BRAND_RED, BRAND_ORANGE][i % 3])
    else:
        for i, pt in enumerate(data.points[:3]):
            parts = re.split(r"[:\–\—]", _clean(pt), 1)
            val = parts[1].strip()[:16] if len(parts) > 1 else pt[:14]
            lbl = parts[0].strip()[:32] if len(parts) > 1 else f"Metric {i+1}"
            _kpi_card(slide, MG + i * (cw_card + cgap), CTOP,
                      cw_card, ch_card, val, lbl,
                      [BRAND_GREEN, BRAND_RED, BRAND_ORANGE][i % 3])

    # Body section
    body_t = CTOP + ch_card + Inches(0.14)
    ins_h = Inches(1.08) if data.insight else Inches(0)
    body_h = CBOT - body_t - ins_h - Inches(0.08)

    lw = Inches(5.40)
    rx = MG + lw + Inches(0.22)
    rw = W - rx - MG

    labels, values = _extract_bar_data(data.points)
    if len(values) >= 2:
        try:
            img = _chart_bar(labels, values,
                             w_in=lw.inches - 0.12, h_in=body_h.inches - 0.10)
            _rrect(slide, MG, body_t, lw, body_h, CARD_BG, border=CARD_BOR, bpt=0.4, adj=2500)
            _txt(slide, "COST / EFFICIENCY BREAKDOWN",
                 MG + Inches(0.15), body_t + Inches(0.07),
                 lw - Inches(0.3), Inches(0.22),
                 size=7.5, bold=True, color=BRAND_GREEN)
            _pic(slide, img, MG + Inches(0.08), body_t + Inches(0.30),
                 lw - Inches(0.16), body_h - Inches(0.38))
        except Exception:
            pass
        _bullets(slide, data.points, rx, body_t, rw, body_h, fs=10, accent=BRAND_GREEN)
    else:
        _bullets(slide, data.points, MG, body_t, CW, body_h, fs=10.5, accent=BRAND_GREEN)

    if data.insight:
        _insight_box(slide, data.insight, MG, CBOT - ins_h, CW)

# ── Slide 7: Risks, Actions & Next Steps ─────────────────────────────────────

def _slide_actions(slide, data: SlideOutline, total: int, event_name: str):
    _header(slide, data.title, data.number, total, BRAND_RED)
    _footer(slide, event_name)

    col_w = (CW - Inches(0.26)) / 2
    rx = MG + col_w + Inches(0.26)

    # Classify points into risks vs actions
    risk_kw = ("risk", "gap", "miss", "lack", "fail", "low", "weak", "poor", "below",
                "limited", "insuff", "no data", "challeng")
    act_kw  = ("action", "next", "priorit", "recommend", "implement", "improve",
               "increase", "focus", "invest", "target", "optimize")
    risks   = [p for p in data.points if any(k in p.lower() for k in risk_kw)] or data.points[:3]
    actions = [p for p in data.points if any(k in p.lower() for k in act_kw)]
    if not actions:
        actions = data.points[len(risks):] or data.points

    avail_h = CBOT - CTOP - Inches(0.32)

    # Left: Risk cards
    _txt(slide, "▸  RISKS & GAPS",
         MG, CTOP, col_w, Inches(0.28),
         size=9, bold=True, color=BRAND_RED)
    n_r = min(len(risks), 4)
    rcard_h = (avail_h - (n_r - 1) * Inches(0.10)) / max(n_r, 1)
    for i, pt in enumerate(risks[:4]):
        cy = CTOP + Inches(0.32) + i * (rcard_h + Inches(0.10))
        _rrect(slide, MG, cy, col_w, rcard_h,
               RGBColor(0xFF, 0xF5, 0xF5), border=BRAND_RED, bpt=0.8, adj=2500)
        _rect(slide, MG, cy, Inches(0.06), rcard_h, BRAND_RED)
        _txt(slide, _fit(pt, 115),
             MG + Inches(0.16), cy + Inches(0.10),
             col_w - Inches(0.24), rcard_h - Inches(0.18),
             size=9.5, color=TXT_DARK, wrap=True, ls=1.10)

    # Right: Action cards
    _txt(slide, "▸  PRIORITY ACTIONS",
         rx, CTOP, col_w, Inches(0.28),
         size=9, bold=True, color=BRAND_GREEN)
    n_a = min(len(actions), 4)
    acard_h = (avail_h - (n_a - 1) * Inches(0.10)) / max(n_a, 1)
    prio_cols = [BRAND_RED, BRAND_ORANGE, BRAND_ORANGE, BRAND_GREEN]
    for i, pt in enumerate(actions[:4]):
        cy = CTOP + Inches(0.32) + i * (acard_h + Inches(0.10))
        _action_card(slide, rx, cy, col_w, acard_h,
                     i + 1, pt, prio_cols[i % len(prio_cols)])

# ── Table Slide ───────────────────────────────────────────────────────────────

def _slide_table(slide, data: SlideOutline, total: int, event_name: str):
    _header(slide, data.title, data.number, total, BRAND_TEAL)
    _footer(slide, event_name)

    rows = (data.table_data or [])[:7]
    cols = list(rows[0].keys()) if rows else []
    if not cols:
        _slide_standard(slide, data, total, event_name)
        return

    nc = len(cols)
    col_w = CW / nc
    row_h = Inches(0.46)
    y = CTOP + Inches(0.06)

    for j, col in enumerate(cols):
        x = MG + j * col_w
        _rrect(slide, x + Inches(0.02), y, col_w - Inches(0.04), row_h,
               HDR_BG, adj=1800)
        _rect(slide, x + Inches(0.02), y, col_w - Inches(0.04), Inches(0.04), BRAND_TEAL)
        _txt(slide, _fit(col.upper(), 20),
             x + Inches(0.12), y + Inches(0.10),
             col_w - Inches(0.20), row_h - Inches(0.12),
             size=9, bold=True, color=TXT_WHITE, align=PP_ALIGN.CENTER)

    for i, row in enumerate(rows):
        y += row_h + Inches(0.03)
        bg = CARD_BG if i % 2 == 0 else SURFACE
        for j, col in enumerate(cols):
            x = MG + j * col_w
            _rrect(slide, x + Inches(0.02), y, col_w - Inches(0.04), row_h,
                   bg, border=CARD_BOR, bpt=0.3, adj=1800)
            val = str(row.get(col, ""))
            is_num = bool(re.match(r"^\$?[\d,\.]+[kKmMbBxX%]?$", val.strip()))
            _txt(slide, _fit(val, 26),
                 x + Inches(0.10), y + Inches(0.10),
                 col_w - Inches(0.18), row_h - Inches(0.12),
                 size=10, bold=is_num, color=TXT_DARK,
                 align=PP_ALIGN.CENTER if is_num else PP_ALIGN.LEFT)

    if data.insight:
        _insight_box(slide, data.insight, MG, CBOT - Inches(1.05), CW)

# ── Standard (fallback) ───────────────────────────────────────────────────────

def _slide_standard(slide, data: SlideOutline, total: int, event_name: str):
    _header(slide, data.title, data.number, total, BRAND_RED)
    _footer(slide, event_name)

    lw = Inches(3.80)
    rx = MG + lw + Inches(0.22)
    rw = W - rx - MG

    headline, lbl, sub = _headline(data.points)
    _metric_panel(slide, MG, CTOP, lw, CAVAIL, headline, lbl, sub, BRAND_RED)

    ins_h = Inches(1.08) if data.insight else Inches(0)
    bh = CAVAIL - ins_h - Inches(0.05)
    _bullets(slide, data.points, rx, CTOP, rw, bh, fs=10.5, accent=BRAND_RED)

    if data.insight:
        _insight_box(slide, data.insight, rx, CBOT - ins_h, rw)

# ── Router ────────────────────────────────────────────────────────────────────

def _build_slide(slide, data: SlideOutline, total: int, event_name: str):
    st = _stype(data.title, bool(data.table_data), data.number)
    {
        "exec":      _slide_exec,
        "pipeline":  _slide_pipeline,
        "funnel":    _slide_funnel,
        "segment":   _slide_segment,
        "platform":  _slide_platform,
        "financial": _slide_financial,
        "actions":   _slide_actions,
        "table":     _slide_table,
    }.get(st, _slide_standard)(slide, data, total, event_name)

# ── Cover Slide ───────────────────────────────────────────────────────────────

def _build_cover(prs: Presentation, event_name: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide, COVER_BG)

    # Left accent stripe
    _rect(slide, 0, 0, Inches(0.55), H, BRAND_RED)

    # Right panel (darker surface for content list)
    right_x = Inches(8.60)
    _rect(slide, right_x, 0, W - right_x, H, RGBColor(0x10, 0x1C, 0x38))
    _rect(slide, right_x, 0, Inches(0.03), H, RGBColor(0x2A, 0x45, 0x75))

    # Eyebrow
    _txt(slide, "POST-EVENT ROI ANALYSIS",
         Inches(0.80), Inches(1.10), Inches(7.5), Inches(0.44),
         size=11, bold=True, color=BRAND_RED)

    # Event name
    name_h = Inches(2.30)
    _txt(slide, event_name,
         Inches(0.80), Inches(1.62), Inches(7.50), name_h,
         size=40, bold=True, color=TXT_WHITE, wrap=True, ls=1.04)

    # Accent divider
    _rect(slide, Inches(0.80), Inches(4.30), Inches(3.80), Inches(0.05), BRAND_RED)

    # Sub-label
    _txt(slide, "Executive Intelligence Report",
         Inches(0.80), Inches(4.46), Inches(7.0), Inches(0.38),
         size=14, italic=True, color=RGBColor(0x7A, 0x93, 0xB8))

    # Contents list on right panel
    _txt(slide, "REPORT CONTENTS",
         right_x + Inches(0.26), Inches(0.72), Inches(3.8), Inches(0.28),
         size=8, bold=True, color=RGBColor(0x6A, 0x85, 0xAA))

    contents = [
        ("01", "Executive Summary"),
        ("02", "Pipeline & Revenue"),
        ("03", "Funnel & Conversion"),
        ("04", "Audience & Segments"),
        ("05", "Platform & Data"),
        ("06", "Financial Efficiency"),
        ("07", "Actions & Next Steps"),
    ]
    for i, (num, title) in enumerate(contents):
        cy = Inches(1.12) + i * Inches(0.64)
        row_bg = RGBColor(0x18, 0x2A, 0x48) if i % 2 == 0 else RGBColor(0x10, 0x1C, 0x38)
        _rect(slide, right_x, cy - Inches(0.06), W - right_x, Inches(0.58), row_bg)
        # Number badge
        _txt(slide, num,
             right_x + Inches(0.16), cy, Inches(0.38), Inches(0.38),
             size=9, bold=True, color=BRAND_RED, align=PP_ALIGN.CENTER)
        _txt(slide, title,
             right_x + Inches(0.60), cy, Inches(3.5), Inches(0.36),
             size=11, color=RGBColor(0xC2, 0xD4, 0xEA))

    # Footer branding
    _txt(slide, "Powered by Gemini AI  ·  Databricks Genie  ·  LatentView MARDI",
         Inches(0.80), H - Inches(0.55), Inches(7.5), Inches(0.36),
         size=9, color=RGBColor(0x3E, 0x55, 0x78))
    _txt(slide, "CONFIDENTIAL",
         W - Inches(1.80), H - Inches(0.44), Inches(1.50), Inches(0.28),
         size=8.5, bold=True, color=RGBColor(0x3E, 0x55, 0x78), align=PP_ALIGN.RIGHT)

# ── Entry Point ───────────────────────────────────────────────────────────────

def generate_pptx(outline: ReportOutline, on_progress=None) -> str:
    """
    Generate a premium PPTX from outline.
    on_progress(slide_num, total_slides, title) called after each slide.
    Returns the saved filename.
    """
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    total = len(outline.slides)

    _build_cover(prs, outline.event_name)
    if on_progress:
        on_progress(0, total + 1, "Cover")

    for sdata in outline.slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _bg(slide, SLIDE_BG)
        _build_slide(slide, sdata, total, outline.event_name)
        if on_progress:
            on_progress(sdata.number, total + 1, sdata.title)

    safe = re.sub(r"[^\w\s\-]", "", outline.event_name).replace(" ", "_")[:40]
    filename = f"roi_{safe}_{uuid.uuid4().hex[:6]}.pptx"
    prs.save(os.path.join(settings.temp_dir, filename))
    return filename
